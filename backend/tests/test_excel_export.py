from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.v1.endpoints import audio_tests
from app.main import app


def test_excel_export_has_requested_columns_and_formatted_comparison_values(monkeypatch):
    async def mock_list_local_runs(limit):
        return [{
            "run_id": "run_123",
            "test_case_id": "STT-123",
            "reference_transcript": "hello world",
            "raw_transcript": "hello there",
            "language_code_input": "en-IN",
            "desired_output": "roman",
            "verbatim": True,
            "translate": False,
            "mode": "verbatim",
            "keyterms": ["hello", "world"],
            "reference_comparison": {"wer": 0.25, "cer_percent": 12.5},
            "external_comparison": {
                "transcript": "hello world",
                "product_vs_reference": {"wer_percent": 20, "cer": 0.3},
            },
        }]

    monkeypatch.setattr(audio_tests, "get_database", lambda: None)
    monkeypatch.setattr(audio_tests, "list_local_runs", mock_list_local_runs)

    response = TestClient(app).get("/api/v1/audio-tests/export.xlsx")

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert 'filename="speechlab-runs.xlsx"' in response.headers["content-disposition"]

    workbook = load_workbook(BytesIO(response.content))
    worksheet = workbook["STT Test Report"]
    assert [cell.value for cell in worksheet[1]] == audio_tests.EXPORT_COLUMNS
    assert [worksheet.cell(row=2, column=column).value for column in range(1, 7)] == [
        "STT-123", "hello world", "hello there", "hello world", "en-IN", "roman"
    ]
    assert [worksheet.cell(row=2, column=column).value for column in (7, 8, 9, 12, 13)] == [
        0.25, 0.125, 0.5, 0.2, 0.3
    ]
    assert worksheet["J2"].value is True
    assert worksheet["K2"].value is False
    assert all(
        worksheet.cell(row=2, column=column).number_format == "0.00%"
        for column in (7, 8, 9, 12, 13)
    )
    assert worksheet.freeze_panes == "A2"
