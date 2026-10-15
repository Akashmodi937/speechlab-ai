import asyncio
import hashlib
import json
import time
import uuid
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Annotated

import av
import httpx
from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.api.v1.endpoints.run_store import (
    list_local_runs,
    save_local_run,
    update_local_run,
)
from app.core.config import settings
from app.db import get_database

router = APIRouter(prefix="/audio-tests", tags=["Audio tests"])
SARVAM_STT_URL = "https://api.sarvam.ai/speech-to-text"
SARVAM_BATCH_JOB_URL = "https://api.sarvam.ai/speech-to-text/job/v1"
SARVAM_REST_MAX_DURATION_SECONDS = 30
SARVAM_BATCH_MAX_DURATION_SECONDS = 2 * 60 * 60
BATCH_POLL_INTERVAL_SECONDS = 5
BATCH_POLL_TIMEOUT_SECONDS = 600
# These codes and options follow the Saaras v4 REST documentation.
LANGUAGE_CODES = {
    "auto": "unknown", "as-IN": "as-IN", "bn-IN": "bn-IN", "brx-IN": "brx-IN",
    "doi-IN": "doi-IN", "en-IN": "en-IN",
    "gu-IN": "gu-IN", "hi-IN": "hi-IN", "kn-IN": "kn-IN", "kok-IN": "kok-IN",
    "ks-IN": "ks-IN", "mai-IN": "mai-IN", "ml-IN": "ml-IN", "mni-IN": "mni-IN",
    "mr-IN": "mr-IN", "ne-IN": "ne-IN", "od-IN": "od-IN", "pa-IN": "pa-IN",
    "sa-IN": "sa-IN", "sat-IN": "sat-IN", "sd-IN": "sd-IN", "ta-IN": "ta-IN",
    "te-IN": "te-IN", "ur-IN": "ur-IN",
}
ALLOWED_MODES = {"transcribe", "translate", "verbatim", "translit", "codemix"}
ALLOWED_EXTENSIONS = {
    ".wav", ".mp3", ".m4a", ".mp4", ".mpeg", ".mpga", ".webm",
    ".ogg", ".flac", ".aac", ".aiff", ".opus", ".amr", ".wma", ".pcm"
}
EXPORT_COLUMNS = [
    "TestId",
    "Reference Transcript",
    "Raw transcript",
    "Whisper + OpenAI Raw Transcript",
    "Input language Mode",
    "Desired Output",
    "WER/Sarvam",
    "CER/Sarvam",
    "Critical term accuracy",
    "Verbatim",
    "Translate",
    "WER/Whisper",
    "CER/Whisper",
]

def now_utc():
    return datetime.now(timezone.utc)

def json_safe(doc: dict) -> dict:
    doc = dict(doc)
    doc.pop("_id", None)
    for key, value in list(doc.items()):
        if isinstance(value, datetime):
            doc[key] = value.isoformat()
    return doc

def get_value(run: dict, *keys: str, default=""):
    for key in keys:
        value = run.get(key)
        if value is not None:
            return value
    return default

def percentage(value):
    if value is None or value == "":
        return ""
    if isinstance(value, str) and "%" in value:
        value = value.split("%", 1)[0].strip()
    try:
        number = float(value)
        return number if number <= 1 else number / 100
    except (ValueError, TypeError):
        return value

def detect_audio_duration_seconds(audio: bytes, filename: str, input_audio_codec: str) -> float:
    if filename.lower().endswith(".pcm"):
        if input_audio_codec == "auto":
            raise ValueError("Raw PCM duration requires an explicit 16 kHz input_audio_codec.")
        if len(audio) % 2:
            raise ValueError("Raw 16 kHz PCM audio must contain complete 16-bit samples.")
        return len(audio) / (16_000 * 2)

    try:
        with av.open(BytesIO(audio), mode="r") as container:
            audio_stream = next(
                (stream for stream in container.streams if stream.type == "audio"),
                None,
            )
            if audio_stream is None:
                raise ValueError("The uploaded file contains no audio stream.")
            if audio_stream.duration is not None and audio_stream.time_base is not None:
                duration = float(audio_stream.duration * audio_stream.time_base)
            elif container.duration is not None:
                duration = container.duration / av.time_base
            else:
                samples = 0
                sample_rate = None
                for frame in container.decode(audio_stream):
                    samples += frame.samples
                    sample_rate = frame.sample_rate or sample_rate
                if not sample_rate:
                    raise ValueError("Could not determine the uploaded audio duration.")
                duration = samples / sample_rate
    except (av.error.FFmpegError, StopIteration) as exc:
        raise ValueError(f"Could not decode uploaded audio to determine its duration: {exc}") from exc

    if duration <= 0:
        raise ValueError("Could not determine the uploaded audio duration.")
    return duration

def _check_batch_response(response: httpx.Response, stage: str) -> None:
    try:
        response.raise_for_status()
    except httpx.HTTPStatusError as exc:
        raise RuntimeError(
            f"Sarvam Batch {stage} failed with HTTP {exc.response.status_code}: "
            f"{exc.response.text[:1000]}"
        ) from exc

async def _transcribe_with_batch(
    client: httpx.AsyncClient,
    audio: bytes,
    filename: str,
    content_type: str,
    job_parameters: dict,
) -> tuple[dict, str]:
    headers = {"api-subscription-key": settings.sarvam_api_key}
    response = await client.post(
        SARVAM_BATCH_JOB_URL,
        headers=headers,
        json={"job_parameters": job_parameters},
    )
    _check_batch_response(response, "job initiation")
    job_id = response.json().get("job_id")
    if not isinstance(job_id, str) or not job_id:
        raise ValueError("Sarvam Batch job initiation returned no job_id.")

    response = await client.post(
        f"{SARVAM_BATCH_JOB_URL}/upload-files",
        headers=headers,
        json={"job_id": job_id, "files": [filename]},
    )
    _check_batch_response(response, "upload URL request")
    upload_response = response.json()
    upload_details = upload_response.get("upload_urls", {}).get(filename, {})
    upload_url = upload_details.get("file_url")
    if not isinstance(upload_url, str) or not upload_url:
        raise ValueError("Sarvam Batch did not return an upload URL for the audio file.")

    upload_headers = dict(upload_details.get("file_metadata") or {})
    storage_container_type = upload_response.get("storage_container_type")
    if storage_container_type in {"Azure", "Azure_V1"} and not any(
        header.lower() == "x-ms-blob-type" for header in upload_headers
    ):
        upload_headers["x-ms-blob-type"] = "BlockBlob"
    if not any(header.lower() == "content-type" for header in upload_headers):
        upload_headers["Content-Type"] = content_type
    response = await client.put(upload_url, content=audio, headers=upload_headers)
    _check_batch_response(response, "audio upload")

    response = await client.post(f"{SARVAM_BATCH_JOB_URL}/{job_id}/start", headers=headers)
    _check_batch_response(response, "job start")

    deadline = asyncio.get_running_loop().time() + BATCH_POLL_TIMEOUT_SECONDS
    while True:
        response = await client.get(f"{SARVAM_BATCH_JOB_URL}/{job_id}/status", headers=headers)
        _check_batch_response(response, "job status")
        status = response.json()
        job_state = status.get("job_state")
        if job_state in {"Completed", "PartiallyCompleted", "Failed"}:
            break
        if job_state not in {"Accepted", "Pending", "Running"}:
            raise ValueError(f"Sarvam Batch returned an unknown job state: {job_state!r}.")
        remaining = deadline - asyncio.get_running_loop().time()
        if remaining <= 0:
            raise TimeoutError(
                f"Sarvam Batch job {job_id} did not complete within "
                f"{BATCH_POLL_TIMEOUT_SECONDS} seconds."
            )
        await asyncio.sleep(min(BATCH_POLL_INTERVAL_SECONDS, remaining))

    if job_state == "Failed":
        raise RuntimeError(status.get("error_message") or f"Sarvam Batch job {job_id} failed.")

    job_details = status.get("job_details") or []
    successful_detail = next(
        (detail for detail in job_details if detail.get("state") == "Success"),
        None,
    )
    if successful_detail is None:
        failure = next(
            (detail for detail in job_details if detail.get("error_message")),
            None,
        )
        reason = failure.get("error_message") if failure else "No successful file result was returned."
        raise RuntimeError(f"Sarvam Batch transcription failed: {reason}")
    output_files = [
        output.get("file_name")
        for output in successful_detail.get("outputs", [])
        if output.get("file_name")
    ]
    if not output_files:
        raise ValueError("Sarvam Batch completed without returning a transcript output file.")

    response = await client.post(
        f"{SARVAM_BATCH_JOB_URL}/download-files",
        headers=headers,
        json={"job_id": job_id, "files": output_files},
    )
    _check_batch_response(response, "download URL request")
    download_urls = response.json().get("download_urls", {})
    result_url = download_urls.get(output_files[0], {}).get("file_url")
    if not isinstance(result_url, str) or not result_url:
        raise ValueError("Sarvam Batch did not return a download URL for its transcript.")

    response = await client.get(result_url)
    _check_batch_response(response, "result download")
    result = response.json()
    if not isinstance(result, dict):
        raise TypeError("Sarvam Batch result must be a JSON object.")
    if not isinstance(result.get("transcript"), str):
        raise TypeError("Sarvam Batch result must contain a string transcript.")
    return result, job_id

def _comparison_metric(run: dict, aliases: tuple[str, ...], comparison: dict, key: str):
    value = get_value(run, *aliases, default=None)
    if value is None:
        value = get_value(comparison, key, f"{key}_percent", default=None)
    return percentage(value)

def create_excel_report(runs: list[dict]) -> BytesIO:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "STT Test Report"
    worksheet.append(EXPORT_COLUMNS)

    widths = [27, 55, 55, 55, 22, 18, 16, 16, 24, 12, 12, 16, 16]
    for run in runs:
        reference_comparison = run.get("reference_comparison") or {}
        external_comparison = run.get("external_comparison") or {}
        whisper_reference = external_comparison.get("product_vs_reference") or {}
        raw_transcript = str(run.get("raw_transcript") or "")
        keyterms = run.get("keyterms") or run.get("critical_terms") or []
        accuracy = get_value(run, "critical_term_accuracy", default=None)
        if accuracy is None and keyterms:
            matched = sum(str(term).casefold() in raw_transcript.casefold() for term in keyterms)
            accuracy = matched / len(keyterms)
        desired_output = get_value(run, "desired_output", "desiredOutput", default=None)
        if desired_output is None:
            desired_output = {
                "codemix": "auto",
                "transcribe": "devanagari",
                "translit": "roman",
            }.get(run.get("mode"), "")
        worksheet.append([
            get_value(run, "test_case_id", "test_id", "testId", "testIdGenerated", "run_id"),
            get_value(run, "reference_transcript", "referenceTranscript"),
            get_value(run, "raw_transcript", "sarvam_raw_transcript"),
            get_value(
                run,
                "whisper_openai_raw_transcript",
                "whisper_transcript",
                "whisper_raw_transcript",
                default=external_comparison.get("transcript", ""),
            ),
            get_value(run, "language_code_input", "input_language_mode", "inputLanguageMode"),
            desired_output,
            _comparison_metric(run, ("wer_sarvam", "sarvam_wer"), reference_comparison, "wer"),
            _comparison_metric(run, ("cer_sarvam", "sarvam_cer"), reference_comparison, "cer"),
            percentage(accuracy),
            get_value(run, "verbatim", default=run.get("mode") == "verbatim"),
            get_value(run, "translate", default=run.get("mode") == "translate"),
            _comparison_metric(run, ("wer_whisper", "whisper_wer"), whisper_reference, "wer"),
            _comparison_metric(run, ("cer_whisper", "whisper_cer"), whisper_reference, "cer"),
        ])

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="315EFB")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in (7, 8, 9, 12, 13):
            worksheet.cell(row=row[0].row, column=column).number_format = "0.00%"

    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 35

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output

@router.post("/run")
async def create_stt_run(
    file: Annotated[UploadFile, File(...)],
    language_code: Annotated[str, Form()] = "auto",
    mode: Annotated[str, Form()] = "transcribe",
    keyterms_json: Annotated[str, Form()] = "[]",
    reference_transcript: Annotated[str, Form()] = "",
    test_case_id: Annotated[str, Form()] = "",
    desired_output: Annotated[str, Form()] = "",
    verbatim: Annotated[bool, Form()] = False,
    translate: Annotated[bool, Form()] = False,
    with_timestamps: Annotated[bool, Form()] = False,
    input_audio_codec: Annotated[str, Form()] = "auto",
):
    if language_code not in LANGUAGE_CODES:
        raise HTTPException(422, "Unsupported language_code for this UI.")
    if mode not in ALLOWED_MODES:
        raise HTTPException(422, f"mode must be one of: {', '.join(sorted(ALLOWED_MODES))}.")
    try:
        keyterms = json.loads(keyterms_json or "[]")
    except json.JSONDecodeError:
        raise HTTPException(422, "keyterms_json must be a JSON array of strings.")
    if not isinstance(keyterms, list) or any(not isinstance(item, str) or not item.strip() for item in keyterms):
        raise HTTPException(422, "keyterms_json must be an array of non-empty strings.")
    if len(keyterms) > 50 or len(set(keyterms)) != len(keyterms) or any(len(term) > 64 for term in keyterms):
        raise HTTPException(422, "Saaras v4 keyterms allow up to 50 distinct terms, max 64 characters each.")
    filename = Path(file.filename or "audio.bin").name
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise HTTPException(415, f"Unsupported audio extension: {suffix}")
    if input_audio_codec not in {"auto", "pcm_s16le", "pcm_l16", "pcm_raw"}:
        raise HTTPException(422, "input_audio_codec must be auto, pcm_s16le, pcm_l16, or pcm_raw.")
    if input_audio_codec != "auto" and suffix != ".pcm":
        raise HTTPException(422, "Choose a .pcm file when setting input_audio_codec; PCM must be 16 kHz.")
    if not settings.sarvam_api_key:
        raise HTTPException(503, "SARVAM_API_KEY is not configured on the backend.")
    audio = await file.read()
    if not audio:
        raise HTTPException(400, "Uploaded audio file is empty.")
    if len(audio) > settings.max_audio_size_mb * 1024 * 1024:
        raise HTTPException(413, f"Audio exceeds {settings.max_audio_size_mb} MB.")
    content_type = file.content_type or "application/octet-stream"
    content_type = content_type.split(";", 1)[0].strip().lower()
    run_id = f"run_{uuid.uuid4().hex}"
    digest = hashlib.sha256(audio).hexdigest()
    storage = Path(settings.audio_storage_dir)
    if not storage.is_absolute():
        storage = (Path(__file__).resolve().parents[5] / storage).resolve()
    storage.mkdir(parents=True, exist_ok=True)
    stored_path = storage / f"{run_id}{suffix}"
    stored_path.write_bytes(audio)
    doc = {
        "run_id": run_id, "created_at": now_utc(), "provider": "sarvam",
        "model": "saaras:v4", "mode": mode, "language_code_input": language_code,
        "language_code_sent": LANGUAGE_CODES[language_code], "keyterms": keyterms,
        "audio_filename": filename, "audio_path": str(stored_path),
        "audio_sha256": digest, "audio_bytes": len(audio),
        "reference_transcript": reference_transcript.strip(), "test_case_id": test_case_id or None,
        "desired_output": desired_output, "verbatim": verbatim, "translate": translate,
        "with_timestamps": with_timestamps, "input_audio_codec": input_audio_codec,
        "status": "running",
    }
    started = time.perf_counter()
    data = {"model": "saaras:v4", "mode": mode, "language_code": LANGUAGE_CODES[language_code]}
    if keyterms:
        data["keyterms"] = json.dumps(keyterms, ensure_ascii=False)
    if with_timestamps:
        data["with_timestamps"] = "true"
    if input_audio_codec != "auto":
        data["input_audio_codec"] = input_audio_codec
    try:
        duration_seconds = detect_audio_duration_seconds(audio, filename, input_audio_codec)
        if duration_seconds > SARVAM_BATCH_MAX_DURATION_SECONDS:
            raise ValueError("Sarvam Batch supports audio up to 2 hours.")
        transcription_api = (
            "batch" if duration_seconds > SARVAM_REST_MAX_DURATION_SECONDS else "rest"
        )
        doc.update({
            "audio_duration_seconds": round(duration_seconds, 3),
            "transcription_api": transcription_api,
        })
        async with httpx.AsyncClient(timeout=httpx.Timeout(120.0, connect=15.0)) as client:
            if transcription_api == "batch":
                batch_parameters = {
                    "model": "saaras:v4",
                    "mode": mode,
                    "language_code": LANGUAGE_CODES[language_code],
                }
                if keyterms:
                    batch_parameters["keyterms"] = keyterms
                if with_timestamps:
                    batch_parameters["with_timestamps"] = True
                if input_audio_codec != "auto":
                    batch_parameters["input_audio_codec"] = input_audio_codec
                result, batch_job_id = await _transcribe_with_batch(
                    client, audio, filename, content_type, batch_parameters
                )
                doc["batch_job_id"] = batch_job_id
            else:
                response = await client.post(
                    SARVAM_STT_URL,
                    headers={"api-subscription-key": settings.sarvam_api_key},
                    data=data,
                    files={"file": (filename, audio, content_type)},
                )
                response.raise_for_status()
                result = response.json()
        doc.update({
            "status": "success", "raw_transcript": result.get("transcript", ""),
            "detected_language": result.get("language_code"),
            "language_probability": result.get("language_probability"),
            "provider_request_id": result.get("request_id"), "timestamps": result.get("timestamps"),
            "latency_seconds": round(time.perf_counter() - started, 3),
        })
    except httpx.HTTPStatusError as exc:
        doc.update({"status": "failed", "latency_seconds": round(time.perf_counter() - started, 3),
                    "error": f"Sarvam HTTP {exc.response.status_code}: {exc.response.text[:1000]}"})
    except Exception as exc:
        doc.update({"status": "failed", "latency_seconds": round(time.perf_counter() - started, 3),
                    "error": f"{type(exc).__name__}: {str(exc)[:1000]}"})
    if doc.get("status") == "success" and doc.get("reference_transcript", "").strip():
        from app.api.v1.endpoints.transcript_tools import compare_text_values
        try:
            doc["reference_comparison"] = compare_text_values(
                doc["reference_transcript"], str(doc.get("raw_transcript") or "")
            )
        except HTTPException:
            doc["reference_comparison"] = None
    db = get_database()
    if db is not None:
        await db["audio_test_runs"].insert_one(doc)
    else:
        await save_local_run(doc)
    return json_safe(doc)

@router.get("/runs")
async def list_runs(limit: int = 50):
    limit = max(1, min(limit, 200))
    db = get_database()
    if db is None:
        return {"mode": "local_file_store", "items": await list_local_runs(limit)}
    cursor = db["audio_test_runs"].find({}, {"_id": 0}).sort("created_at", -1).limit(limit)
    return {"mode": "mongodb", "items": [json_safe(item) async for item in cursor]}

@router.post("/import-external")
async def import_external_result(payload: dict):
    transcript = payload.get("transcript")
    if not isinstance(transcript, str) or not transcript.strip():
        raise HTTPException(422, "JSON body must include a non-empty transcript.")
    doc = {
        "run_id": f"external_{uuid.uuid4().hex}", "created_at": now_utc(),
        "provider": str(payload.get("provider", "external")), "model": payload.get("model"),
        "status": "imported", "raw_transcript": transcript,
        "language_code_input": payload.get("language"), "audio_sha256": payload.get("audio_sha256"),
        "latency_seconds": payload.get("latency_seconds"), "external_metadata": payload.get("metadata", {}),
    }
    db = get_database()
    if db is not None:
        await db["audio_test_runs"].insert_one(doc)
    else:
        await save_local_run(doc)
    return json_safe(doc)

@router.post("/{run_id}/compare-external")
async def compare_external_for_run(run_id: str, payload: dict):
    """Attach a pasted third-party transcript to one Sarvam run and compare the same audio's outputs."""
    transcript = payload.get("transcript")
    if not isinstance(transcript, str) or not transcript.strip():
        raise HTTPException(422, "transcript must be a non-empty string.")
    if len(transcript) > 30000:
        raise HTTPException(422, "transcript must be 30,000 characters or fewer.")
    db = get_database()
    if db is not None:
        run = await db["audio_test_runs"].find_one({"run_id": run_id}, {"_id": 0})
    else:
        run = next((item for item in await list_local_runs(10000) if item.get("run_id") == run_id), None)
    if not run:
        raise HTTPException(404, "Run not found. Run the audio through Saaras first, then paste the other product's result here.")
    sarvam_text = str(run.get("raw_transcript") or "")
    if not sarvam_text.strip():
        raise HTTPException(409, "This Sarvam run has no transcript to compare (it may have failed).")
    from app.api.v1.endpoints.transcript_tools import compare_text_values
    metrics = compare_text_values(sarvam_text, transcript)
    reference_text = str(run.get("reference_transcript") or "").strip()
    external = {"provider": str(payload.get("provider") or "External product")[:120],
                "model": str(payload.get("model") or "")[:120], "transcript": transcript.strip(),
                "sarvam_transcript": sarvam_text, "comparison": metrics,
                "sarvam_vs_reference": run.get("reference_comparison"),
                "product_vs_reference": compare_text_values(reference_text, transcript) if reference_text else None,
                "compared_at": now_utc().isoformat()}
    run["external_comparison"] = external
    if db is not None:
        await db["audio_test_runs"].update_one({"run_id": run_id}, {"$set": {"external_comparison": external}})
    else:
        await update_local_run(run_id, {"external_comparison": external})
    return {"run_id": run_id, "sarvam_transcript": sarvam_text, **external}

@router.get("/export.csv")
async def export_runs_csv():
    import csv
    import io
    db = get_database()
    fields = ["run_id", "created_at", "provider", "model", "mode", "language_code_input", "language_code_sent",
              "detected_language", "with_timestamps", "input_audio_codec", "status", "latency_seconds", "audio_filename", "audio_sha256",
              "reference_transcript", "reference_comparison", "raw_transcript", "external_comparison", "error", "keyterms"]
    buffer = io.StringIO(newline="")
    writer = csv.DictWriter(buffer, fieldnames=fields, extrasaction="ignore")
    writer.writeheader()
    if db is not None:
        cursor = db["audio_test_runs"].find({}, {"_id": 0}).sort("created_at", -1)
        async for item in cursor:
            row = json_safe(item)
            row["keyterms"] = json.dumps(row.get("keyterms", []), ensure_ascii=False)
            row["reference_comparison"] = json.dumps(row.get("reference_comparison", {}), ensure_ascii=False)
            row["external_comparison"] = json.dumps(row.get("external_comparison", {}), ensure_ascii=False)
            writer.writerow({key: row.get(key, "") for key in fields})
    else:
        for item in await list_local_runs(10000):
            row = json_safe(item)
            row["keyterms"] = json.dumps(row.get("keyterms", []), ensure_ascii=False)
            row["reference_comparison"] = json.dumps(row.get("reference_comparison", {}), ensure_ascii=False)
            row["external_comparison"] = json.dumps(row.get("external_comparison", {}), ensure_ascii=False)
            writer.writerow({key: row.get(key, "") for key in fields})
    content = "\ufeff" + buffer.getvalue()
    return StreamingResponse(iter([content.encode("utf-8")]), media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="speechlab-runs.csv"'})

@router.get("/export.xlsx")
async def export_runs_excel():
    db = get_database()
    if db is not None:
        cursor = db["audio_test_runs"].find({}, {"_id": 0}).sort("created_at", -1)
        runs = [json_safe(item) async for item in cursor]
    else:
        runs = await list_local_runs(10000)
    output = create_excel_report(runs)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="speechlab-runs.xlsx"'},
    )
